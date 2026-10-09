"""Lead finder: what a lead list holds, how it is scored, and what may be stored.

A lead is an organisation, never a person (docs/specs/2026-10-07-lead-intelligence-design.md).
SIP keeps company facts with the page each fact came from, generic contact details only,
and a link to the company's LinkedIn people page so sales can find the right person there.

The model reads facts from pages; the score is computed here from the scorecard the user
agreed, so the same facts always give the same score and the scorecard can change without
searching again.
"""

from __future__ import annotations

from io import BytesIO
import re
from typing import Literal
from urllib.parse import urlparse

from pydantic import BaseModel, Field, model_validator

MAX_LEADS = 50
BATCH_SIZE = 10
RETENTION_DAYS = 90
MAX_EXTRA_COLUMNS = 3
MAX_CRITERIA = 4
# Fixed columns a scorecard may use besides the extra columns.
SCORABLE_COLUMNS = ("country", "city")


# --------------------------------------------------------------------------- model output


class LeadColumn(BaseModel):
    name: str  # header the user sees, e.g. "Vestigingen"
    description: str  # what the agent looks for on the website


class ScoreCriterion(BaseModel):
    column: str  # "country", "city" or the name of an extra column
    kind: Literal["number", "text"]
    # number: the minimum for high / medium. text: comma-separated words to look for.
    high: str
    medium: str


class LeadBrief(BaseModel):
    """The search brief the intake produces. `count` stays null until the user names it."""

    title: str
    description: str
    industries: list[str]
    regions: list[str]
    size: str  # e.g. "at least 5 branches", "more than 50,000 inhabitants"
    exclude: list[str]  # what must not be on the list, e.g. "one-man businesses"
    extra_columns: list[LeadColumn]
    scorecard: list[ScoreCriterion]
    count: int | None
    scoring: Literal["levels", "sam_points"] = "levels"

    @model_validator(mode="before")
    @classmethod
    def _stored_before_size_and_exclude(cls, data):
        # Lists and chats saved before these fields existed still load. Not a schema
        # default: the model must always answer both fields.
        if isinstance(data, dict):
            data.setdefault("size", "")
            data.setdefault("exclude", [])
        return data


class LeadIntakeTurn(BaseModel):
    message: str
    brief: LeadBrief
    context_id: str | None = None
    can_search: bool = True
    # True only once the user confirmed the scorecard and named a number of leads.
    ready: bool


class LeadQuery(BaseModel):
    query: str
    country: str  # two-letter code for the search engine, e.g. "nl"


class LeadQueryPlan(BaseModel):
    queries: list[LeadQuery]


class LeadCandidate(BaseModel):
    name: str
    website: str


class LeadCandidates(BaseModel):
    candidates: list[LeadCandidate]


class LeadFact(BaseModel):
    value: str | None = None
    source: str | None = None


class LeadExtraValue(BaseModel):
    column: str
    value: str | None
    source: str | None


class LeadExtraction(BaseModel):
    fits: bool
    name: str
    city: LeadFact
    country: LeadFact
    phone: LeadFact
    email: LeadFact
    extra: list[LeadExtraValue]
    why_fits: str


# --------------------------------------------------------------------------- stored rows


class LeadRow(BaseModel):
    name: str
    website: str
    city: LeadFact = Field(default_factory=LeadFact)
    country: LeadFact = Field(default_factory=LeadFact)
    phone: LeadFact = Field(default_factory=LeadFact)
    email: LeadFact = Field(default_factory=LeadFact)
    linkedin: str | None = None
    extra: dict[str, LeadFact] = Field(default_factory=dict)
    why_fits: str = ""
    sources: list[str] = Field(default_factory=list)
    signals: dict[str, LeadFact] = Field(default_factory=dict)


class CriterionScore(BaseModel):
    column: str
    value: str | None
    level: Literal["high", "medium", "low", "unknown"]


class LeadScore(BaseModel):
    level: Literal["high", "medium", "low"] | None
    criteria: list[CriterionScore]
    points: float | None = None
    tier: str | None = None
    blocks: dict[str, float] = Field(default_factory=dict)
    components: dict[str, float] = Field(default_factory=dict)
    unknown: list[str] = Field(default_factory=list)


SAM_COLUMNS = [
    LeadColumn(name="Vestigingen", description="Aantal vestigingen van de dealergroep; tel geen afdelingen. Alleen expliciet onderbouwde aantallen."),
    LeadColumn(name="Aantal merken", description="Aantal verschillende automerken van de groep, onderbouwd op de eigen website."),
    LeadColumn(name="Open in het weekend", description="Ja als klantenservice/service in het weekend open is, nee alleen bij expliciet gesloten zaterdag EN zondag; anders onbekend."),
]
SAM_SIGNALS = {"rating": "Rating", "reviews": "Reviews", "vacancies": "Vacatures", "chat": "Chattool"}
SAM_EXPLANATION = "Schaal 30: vestigingen × 2 (max 20) + reviews / 250 (max 10). Pijn 30: rating <3,8 = 12; <4,3 = 6; <4,6 = 3; anders 0. Klachten (18) worden niet verzameld: reviewteksten zijn uitgesloten in afwachting van privacybeoordeling. Koopsignaal 25: receptie 12, klantenservice 10, serviceadviseur 7; tel aanwezige categorieën op + 3 bij minstens één vacature, max 25. Fit 15: geen chattool gedetecteerd 6, service weekend gesloten 4, minstens 4 merken 5. A ≥65, B ≥45, C <45. Maximaal haalbaar 82 van 100. Onbekende feiten leveren 0 punten op. Geen chat gedetecteerd is een HTML-signaal, geen bewijs van afwezigheid. Google-cijfers gelden alleen voor gevonden, op domein gematchte vestigingen."


# --------------------------------------------------------------------------- brief


def clean_brief(brief: LeadBrief) -> LeadBrief:
    """What SIP accepts, whatever the model or the browser sent (fail closed)."""
    columns: list[LeadColumn] = []
    for column in brief.extra_columns:
        name = " ".join(column.name.split())[:40]
        if name and name.casefold() not in {c.name.casefold() for c in columns} and name.casefold() not in SCORABLE_COLUMNS:
            columns.append(LeadColumn(name=name, description=column.description.strip()[:300]))
    columns = SAM_COLUMNS if brief.scoring == "sam_points" else columns[:MAX_EXTRA_COLUMNS]
    allowed = {*SCORABLE_COLUMNS, *(c.name.casefold() for c in columns)}
    criteria = [
        ScoreCriterion(column=_column_key(c.column, columns), kind=c.kind, high=c.high.strip()[:200], medium=c.medium.strip()[:200])
        for c in brief.scorecard
        if c.column.strip().casefold() in allowed
    ][:MAX_CRITERIA]
    if brief.scoring == "sam_points":
        criteria = []
    count = None if brief.count is None else max(1, min(MAX_LEADS, brief.count))
    return LeadBrief(
        title=" ".join(brief.title.split())[:120] or "Lead list",
        description=brief.description.strip()[:1000],
        industries=[item.strip()[:80] for item in brief.industries if item.strip()][:8],
        regions=[item.strip()[:80] for item in brief.regions if item.strip()][:8],
        size=brief.size.strip()[:300],
        exclude=[item.strip()[:120] for item in brief.exclude if item.strip()][:8],
        extra_columns=columns,
        scorecard=criteria,
        count=count,
        scoring=brief.scoring,
    )


MIN_USER_TURNS = 2


def brief_complete(brief: LeadBrief) -> bool:
    """Everything the search needs. What to exclude may be empty: "nothing" is an answer."""
    return bool(brief.description and brief.industries and brief.regions and brief.size and (brief.scoring == "sam_points" or brief.scorecard) and brief.count)


def _column_key(name: str, columns: list[LeadColumn]) -> str:
    """The canonical spelling: fixed columns lower-case, extra columns as the user wrote them."""
    folded = name.strip().casefold()
    if folded in SCORABLE_COLUMNS:
        return folded
    return next(column.name for column in columns if column.name.casefold() == folded)


# --------------------------------------------------------------------------- score

POINTS = {"high": 2, "medium": 1, "low": 0, "unknown": 0}


def _number(text: str) -> float | None:
    match = re.search(r"\d[\d.,]*", text)
    if not match:
        return None
    raw = match.group(0).rstrip(".,")
    # "1.200" and "1,200" are thousands; "2,5" is a decimal.
    if re.fullmatch(r"\d{1,3}([.,]\d{3})+", raw):
        raw = re.sub(r"[.,]", "", raw)
    try:
        return float(raw.replace(",", "."))
    except ValueError:
        return None


def _words(text: str) -> list[str]:
    return [word.strip().casefold() for word in text.split(",") if word.strip()]


def _positive_keyword_match(keywords: str, value: str) -> bool:
    """Literal word evidence, allowing intervening words but no negated clauses.

    Comma-separated alternatives remain alternatives. Whole words avoid matching
    'gas' in 'gasten'; every word in a multiword criterion must be present.
    """
    clauses = re.split(r"[,.;\n]|\b(?:maar|but|aber)\b", value.casefold())
    negative = re.compile(r"\b(?:geen|niet|no|not|without|kein|keine|keinen|nicht|ohne)\b")
    for alternative in _words(keywords):
        wanted = set(re.findall(r"\w+", alternative))
        if not wanted:
            continue
        for clause in clauses:
            if negative.search(clause) and not negative.search(alternative):
                continue
            if wanted.issubset(set(re.findall(r"\w+", clause))):
                return True
    return False


def _criterion_level(criterion: ScoreCriterion, value: str | None) -> str:
    if not value:
        return "unknown"
    if criterion.kind == "number":
        number = _number(value)
        if number is None:
            return "unknown"
        high, medium = _number(criterion.high), _number(criterion.medium)
        if high is not None and number >= high:
            return "high"
        if medium is not None and number >= medium:
            return "medium"
        return "low"
    if _positive_keyword_match(criterion.high, value):
        return "high"
    if _positive_keyword_match(criterion.medium, value):
        return "medium"
    return "low"


def fact_for(row: LeadRow, column: str) -> LeadFact:
    if column in SCORABLE_COLUMNS:
        return getattr(row, column)
    return row.extra.get(column) or LeadFact()


def score_row(row: LeadRow, scorecard: list[ScoreCriterion], scoring: str = "levels") -> LeadScore:
    """High / medium / low from the scorecard. An unknown fact can never give high.

    High means mostly high: with two criteria both must be high, with three two high
    and one medium. An even mix of high and medium is medium.
    """
    if scoring == "sam_points":
        return sam_score(row)
    criteria = [
        CriterionScore(column=c.column, value=fact_for(row, c.column).value, level=_criterion_level(c, fact_for(row, c.column).value))
        for c in scorecard
    ]
    if not criteria:
        return LeadScore(level=None, criteria=[])
    average = sum(POINTS[item.level] for item in criteria) / len(criteria)
    has_unknown = any(item.level == "unknown" for item in criteria)
    if average > 1.5 and not has_unknown:
        level = "high"
    elif average >= 0.75:
        level = "medium"
    else:
        level = "low"
    return LeadScore(level=level, criteria=criteria)


def sam_score(row: LeadRow) -> LeadScore:
    unknown = []
    def value(key, extra=False):
        fact = (row.extra if extra else row.signals).get(key, LeadFact())
        if not fact.value or not fact.source:
            unknown.append(key)
            return None
        return fact.value
    def number(key, extra=False):
        raw = value(key, extra)
        result = _number(raw) if raw else None
        if raw and result is None:
            unknown.append(key)
        return result
    branches, reviews, rating = number("Vestigingen", True), number("reviews"), number("rating")
    brands = number("Aantal merken", True)
    vacancies, chat, weekend = value("vacancies"), value("chat"), value("Open in het weekend", True)
    categories = set((vacancies or "").split(", "))
    vacancy_points = min(25, sum(p for name, p in (("receptie", 12), ("klantenservice", 10), ("serviceadviseur", 7)) if name in categories) + (3 if categories.intersection({"receptie", "klantenservice", "serviceadviseur"}) else 0))
    components = {
        "Vestigingen": min(20, max(0, branches or 0) * 2),
        "Reviews": min(10, max(0, reviews or 0) // 250),
        "Rating": (12 if rating < 3.8 else 6 if rating < 4.3 else 3 if rating < 4.6 else 0) if rating is not None and 1 <= rating <= 5 else 0,
        "Klachten (uitgesloten)": 0,
        "Vacatures": vacancy_points,
        "Chattool": 6 if chat == "Niet gedetecteerd" else 0,
        "Weekend gesloten": 4 if (weekend or "").strip().casefold() in {"nee", "no", "nein"} else 0,
        "Aantal merken": 5 if brands is not None and brands >= 4 else 0,
    }
    blocks = {"Schaal": components["Vestigingen"] + components["Reviews"], "Pijn": components["Rating"], "Koopsignaal": vacancy_points, "Fit": components["Chattool"] + components["Weekend gesloten"] + components["Aantal merken"]}
    points = round(sum(blocks.values()), 2)
    tier = "A" if points >= 65 else "B" if points >= 45 else "C"
    return LeadScore(level={"A": "high", "B": "medium", "C": "low"}[tier], criteria=[], points=points, tier=tier, blocks={k: round(v, 2) for k, v in blocks.items()}, components={k: round(v, 2) for k, v in components.items()}, unknown=unknown)


# --------------------------------------------------------------------------- contact rules

# Mailbox names that belong to a role or a place, not to a person.
ROLE_WORDS = {
    "info", "contact", "sales", "verkoop", "vente", "verkauf", "office", "kantoor", "admin",
    "administratie", "administration", "hr", "jobs", "vacatures", "directie", "management",
    "receptie", "reception", "service", "klantenservice", "support", "hallo", "hello", "welkom",
    "mail", "post", "finance", "boekhouding", "marketing", "communicatie", "crm", "aftersales",
    "onderdelen", "parts", "werkplaats", "planning", "algemeen", "general", "team", "secretariaat",
}
EMAIL = re.compile(r"^[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}$")
# firstname.lastname, j.peeters, jan_peeters: looks like a person.
PERSONAL_LOCAL = re.compile(r"^[a-z]+[._][a-z]+$")


def generic_email(value: str | None) -> str | None:
    """A role address such as info@ or sales@, or None when it looks like a person's."""
    if not value:
        return None
    value = value.strip().removeprefix("mailto:").split("?")[0].strip()
    if not EMAIL.match(value):
        return None
    local = value.split("@", 1)[0].casefold()
    if PERSONAL_LOCAL.match(local) and not ROLE_WORDS.intersection(re.split(r"[._]", local)):
        return None
    return value


MOBILE_PREFIXES = (
    "316",  # Netherlands, 06
    "3245", "3246", "3247", "3248", "3249",  # Belgium, 04xx
    "4915", "4916", "4917",  # Germany, 015x-017x
)


def _international_digits(value: str) -> str:
    digits = re.sub(r"\D", "", value.replace("+", "00", 1) if value.strip().startswith("+") else value)
    if digits.startswith("00"):
        return digits[2:]
    return digits


def company_phone(value: str | None, country: str | None = None) -> str | None:
    """The company's main number, or None for a mobile number (likely a person's)."""
    if not value:
        return None
    value = value.strip().removeprefix("tel:").strip()
    digits = re.sub(r"\D", "", value)
    if len(digits) < 8 or len(digits) > 15:
        return None
    international = _international_digits(value)
    if international.startswith(MOBILE_PREFIXES):
        return None
    local = digits if not value.strip().startswith("+") else ""
    country = (country or "").casefold()
    if local.startswith("06") and country in ("", "nl", "netherlands", "nederland"):
        return None
    if re.match(r"^04[5-9]", local) and country in ("be", "belgium", "belgië", "belgie"):
        return None
    if re.match(r"^01[5-7]", local) and country in ("de", "germany", "duitsland", "deutschland"):
        return None
    return value


LINKEDIN_COMPANY = re.compile(r"https?://(?:[a-z]{2,3}\.)?linkedin\.com/company/([^/?#\s]+)", re.IGNORECASE)


def linkedin_people_page(url: str | None) -> str | None:
    """The company's LinkedIn people page. A personal profile (/in/...) is never kept."""
    match = LINKEDIN_COMPANY.match(url or "")
    return f"https://www.linkedin.com/company/{match.group(1)}/people/" if match else None


def domain_of(url: str) -> str:
    host = (urlparse(url if "://" in url else f"https://{url}").hostname or "").casefold()
    return host.removeprefix("www.")


def _digits(text: str) -> str:
    return re.sub(r"\D", "", text)


def build_row(candidate: LeadCandidate, pages: dict[str, str], extraction: LeadExtraction, columns: list[LeadColumn]) -> LeadRow:
    """Keep only what the fetched pages back up: a fact without a fetched source is unknown."""
    folded_pages = {url: text.casefold() for url, text in pages.items()}

    def grounded(fact: LeadFact) -> LeadFact:
        if not fact.value or fact.source not in pages:
            return LeadFact()
        return LeadFact(value=" ".join(fact.value.split())[:200], source=fact.source)

    def on_page(fact: LeadFact) -> LeadFact:
        # A place name is copied, not derived: it must appear on the page it came from.
        return fact if fact.value and fact.value.casefold() in folded_pages[fact.source] else LeadFact()

    country = grounded(extraction.country)
    phone = grounded(extraction.phone)
    phone_value = company_phone(phone.value, country.value)
    # The number itself must be on the page it came from; the model may not compose one.
    if phone_value and _digits(phone_value)[-8:] not in _digits(pages[phone.source]):
        phone_value = None
    email = grounded(extraction.email)
    email_value = generic_email(email.value)
    if email_value and email_value.casefold() not in folded_pages[email.source]:
        email_value = None
    by_name = {item.column.strip().casefold(): item for item in extraction.extra}
    extra: dict[str, LeadFact] = {}
    for column in columns:
        item = by_name.get(column.name.casefold())
        extra[column.name] = grounded(LeadFact(value=item.value, source=item.source)) if item else LeadFact()
    row = LeadRow(
        name=" ".join(extraction.name.split())[:160] or candidate.name,
        website=next(iter(pages)),
        city=on_page(grounded(extraction.city)),
        country=country,
        phone=LeadFact(value=phone_value, source=phone.source) if phone_value else LeadFact(),
        email=LeadFact(value=email_value, source=email.source) if email_value else LeadFact(),
        extra=extra,
        why_fits=" ".join(extraction.why_fits.split())[:400],
    )
    used = [row.city, row.country, row.phone, row.email, *row.extra.values()]
    row.sources = list(dict.fromkeys([row.website, *(fact.source for fact in used if fact.source)]))
    return row


# --------------------------------------------------------------------------- export

EXPORT_LABELS = {
    "nl": {
        "score": "Score", "name": "Bedrijfsnaam", "city": "Plaats", "country": "Land", "website": "Website",
        "phone": "Telefoon", "email": "E-mail", "linkedin": "LinkedIn personen", "why": "Waarom past dit",
        "sources": "Bronnen", "open": "openen", "high": "Hoog", "medium": "Midden", "low": "Laag",
        "sheet": "Leads", "about": "Over deze lijst", "context": "Business Context", "created": "Gemaakt op",
        "expires": "Wordt verwijderd op", "scorecard": "Scorekaart", "note": "Alleen organisatiegegevens en algemene contactgegevens. Gevonden op openbare websites; controleer voor gebruik.",
    },
    "en": {
        "score": "Score", "name": "Company", "city": "City", "country": "Country", "website": "Website",
        "phone": "Phone", "email": "Email", "linkedin": "LinkedIn people", "why": "Why this fits",
        "sources": "Sources", "open": "open", "high": "High", "medium": "Medium", "low": "Low",
        "sheet": "Leads", "about": "About this list", "context": "Business Context", "created": "Created",
        "expires": "Deleted on", "scorecard": "Scorecard", "note": "Organisation data and generic contact details only. Found on public websites; check before use.",
    },
    "de": {
        "score": "Score", "name": "Unternehmen", "city": "Ort", "country": "Land", "website": "Website",
        "phone": "Telefon", "email": "E-Mail", "linkedin": "LinkedIn Personen", "why": "Warum es passt",
        "sources": "Quellen", "open": "öffnen", "high": "Hoch", "medium": "Mittel", "low": "Niedrig",
        "sheet": "Leads", "about": "Über diese Liste", "context": "Business Context", "created": "Erstellt am",
        "expires": "Gelöscht am", "scorecard": "Scorecard", "note": "Nur Unternehmensdaten und allgemeine Kontaktdaten. Auf öffentlichen Websites gefunden; vor Gebrauch prüfen.",
    },
}


def export_xlsx(
    brief: LeadBrief,
    rows: list[LeadRow],
    *,
    context_name: str,
    created_at: str,
    expires_at: str,
    language: str = "nl",
) -> bytes:
    """The list as an Excel file in the shape sales already uses (the SAM list)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    labels = EXPORT_LABELS.get(language, EXPORT_LABELS["nl"])
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = labels["sheet"]

    def put(cell, value) -> None:
        # Text only: a company name such as "=HYPERLINK(...)" must never become a formula.
        cell.value = value
        if isinstance(value, str):
            cell.data_type = "s"

    headers = [
        labels["score"], labels["name"], labels["city"], labels["country"], labels["website"], labels["phone"],
        labels["email"], labels["linkedin"], *(column.name for column in brief.extra_columns), *(list(SAM_SIGNALS.values()) + ["Punten", "Tier"] if brief.scoring == "sam_points" else []), labels["why"], labels["sources"],
    ]
    for index, header in enumerate(headers, start=1):
        put(sheet.cell(row=1, column=index), header)
        sheet.cell(row=1, column=index).font = Font(bold=True)
    scored = sorted(
        ((score_row(row, brief.scorecard, brief.scoring), row) for row in rows),
        key=lambda pair: -(pair[0].points if pair[0].points is not None else POINTS.get(pair[0].level or "low", 0)),
    )
    for row_number, (score, row) in enumerate(scored, start=2):
        values = [
            labels[score.level] if score.level else "",
            row.name, row.city.value, row.country.value, row.website, row.phone.value, row.email.value,
            labels["open"] if row.linkedin else None,
            *(row.extra.get(column.name, LeadFact()).value for column in brief.extra_columns),
            *([row.signals.get(key, LeadFact()).value for key in SAM_SIGNALS] + [score.points, score.tier] if brief.scoring == "sam_points" else []),
            row.why_fits, "\n".join(row.sources),
        ]
        for column_number, value in enumerate(values, start=1):
            put(sheet.cell(row=row_number, column=column_number), value)
        if row.linkedin:
            link = sheet.cell(row=row_number, column=8)
            link.hyperlink = row.linkedin
            link.font = Font(color="0563C1", underline="single")
    widths = [9, 34, 18, 10, 30, 16, 28, 16, *([16] * len(brief.extra_columns)), *([22] * 6 if brief.scoring == "sam_points" else []), 60, 50]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "B2"

    about = workbook.create_sheet(labels["about"])
    lines = [
        (labels["context"], context_name),
        (labels["created"], created_at[:10]),
        (labels["expires"], expires_at[:10]),
        ("", ""),
        (labels["scorecard"], ""),
        *(( ("SAM", SAM_EXPLANATION),) if brief.scoring == "sam_points" else ()),
        *((criterion.column, f"{labels['high']}: {criterion.high} · {labels['medium']}: {criterion.medium}") for criterion in brief.scorecard),
        ("", ""),
        (labels["note"], ""),
    ]
    for row_number, (key, value) in enumerate(lines, start=1):
        put(about.cell(row=row_number, column=1), key)
        put(about.cell(row=row_number, column=2), value)
    about.column_dimensions["A"].width = 28
    about.column_dimensions["B"].width = 70

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
