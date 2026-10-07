"""Compare a Lead finder export with a list sales made by hand (the "gold standard").

Step 1 of docs/specs/2026-10-07-lead-intelligence-design.md: does the search find the
same organisations, and do the facts match? Organisations are matched on their website
domain. Both files stay local; neither belongs in the repository.

    python dify/eval_leads.py <export.xlsx> <gold.xlsx> [--column Vestigingen]

The gold file needs the columns Bedrijfsnaam and Website, and optionally Telefoon,
E-mail and the extra column to compare (for the SAM list: Vestigingen).
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "backend"))
from app.leads import domain_of  # noqa: E402


def rows(path: Path) -> list[dict[str, str]]:
    sheet = load_workbook(path, data_only=True).worksheets[0]
    values = list(sheet.iter_rows(values_only=True))
    headers = [str(cell or "").strip() for cell in values[0]]
    return [{header: "" if cell is None else str(cell).strip() for header, cell in zip(headers, row)} for row in values[1:] if any(row)]


def by_domain(items: list[dict[str, str]]) -> dict[str, dict[str, str]]:
    return {domain_of(item.get("Website", "")): item for item in items if item.get("Website")}


def digits(value: str) -> str:
    return re.sub(r"\D", "", value)[-9:]


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("export", type=Path)
    parser.add_argument("gold", type=Path)
    parser.add_argument("--column", default="Vestigingen", help="extra column to compare")
    args = parser.parse_args()

    found, gold = by_domain(rows(args.export)), by_domain(rows(args.gold))
    both = sorted(set(found) & set(gold))
    print(f"Gold standard: {len(gold)} organisations; export: {len(found)}; found in both: {len(both)}")
    checks = {
        "Telefoon": lambda a, b: digits(a) == digits(b),
        "E-mail": lambda a, b: a.casefold() == b.casefold(),
        args.column: lambda a, b: re.sub(r"\D", "", a) == re.sub(r"\D", "", b),
    }
    for column, same in checks.items():
        compared = [(found[d].get(column, ""), gold[d].get(column, "")) for d in both if gold[d].get(column)]
        filled = [(a, b) for a, b in compared if a]
        right = sum(1 for a, b in filled if same(a, b))
        print(f"{column}: {right} of {len(filled)} filled values match; {len(compared) - len(filled)} left empty")
    missing = sorted(set(gold) - set(found))
    if missing:
        print("Not found:", ", ".join(missing))


if __name__ == "__main__":
    main()
