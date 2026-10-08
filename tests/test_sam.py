"""SAM scoring and sourced signals, with no external credentials or requests."""
from io import BytesIO

import httpx
import pytest
from openpyxl import load_workbook

from app.leads import LeadBrief, LeadFact, LeadRow, clean_brief, brief_complete, score_row, export_xlsx
from app.lead_search import Serper, places_signals, detect_chat, classify_vacancy, LeadSearch


def fact(value):
    return LeadFact(value=str(value), source="https://dealer.example/")


def brief():
    return clean_brief(LeadBrief(title="SAM", description="Dealergroepen", industries=["Automotive"], regions=["Limburg"], size="2+ vestigingen", exclude=[], extra_columns=[], scorecard=[], count=3, scoring="sam_points"))


def lead(**signals):
    return LeadRow(name="Dealer", website="https://dealer.example/", extra={"Vestigingen": fact(9), "Aantal merken": fact(7), "Open in het weekend": fact("Nee")}, signals={key: fact(value) for key, value in signals.items()})


def test_sam_brief_has_fixed_columns_and_no_custom_criteria_requirement():
    assert [column.name for column in brief().extra_columns] == ["Vestigingen", "Aantal merken", "Open in het weekend"]
    assert brief_complete(brief())
    assert LeadBrief.model_validate({**brief().model_dump(exclude={"scoring"})}).scoring == "levels"


def test_reference_without_review_complaints_is_sixty():
    score = score_row(lead(rating="4.29", reviews="2.002", vacancies="klantenservice", chat="Niet gedetecteerd"), [], "sam_points")
    assert score.points == 60
    assert score.tier == "B"
    assert score.blocks == {"Schaal": 26, "Pijn": 6, "Koopsignaal": 13, "Fit": 15}


def test_maximum_is_82_and_no_complaints_can_add_points():
    row = lead(rating=3.7, reviews=3000, vacancies="receptie, klantenservice, serviceadviseur", chat="Niet gedetecteerd", complaints=100)
    row.extra["Vestigingen"] = fact(10)
    score = score_row(row, [], "sam_points")
    assert (score.points, score.tier, score.level) == (82, "A", "high")


@pytest.mark.parametrize("rating,points", [(3.79,12), (3.8,6), (4.29,6), (4.3,3), (4.59,3), (4.6,0), (6,0)])
def test_rating_boundaries(rating, points):
    assert score_row(lead(rating=rating), [], "sam_points").components["Rating"] == points


def test_unknown_and_unsourced_are_not_positive_signals():
    row = LeadRow(name="Unknown", website="https://dealer.example/")
    row.signals["chat"] = LeadFact(value="Niet gedetecteerd")
    score = score_row(row, [], "sam_points")
    assert score.points == 0
    assert "chat" in score.unknown
    assert "rating" in score.unknown


def test_places_weight_by_review_count_deduplicate_and_match_exact_domain():
    places = [dict(website="https://dealer.example/a", cid="1", ratingCount=100, rating=3), dict(website="https://www.dealer.example/b", cid="2", ratingCount=300, rating=5), dict(website="https://dealer.example/a", cid="1", ratingCount=100, rating=3), dict(website="https://dealer.example.evil/", cid="3", ratingCount=9000, rating=1)]
    signals = places_signals(places, "https://dealer.example/")
    assert signals["reviews"].value == "400"
    assert float(signals["rating"].value) == 4.5
    assert signals["place_2"].source.endswith("cid=2")
    assert places_signals([], "https://dealer.example/") == {}


def test_serper_drops_personal_fields_and_never_requests_reviews():
    def respond(request):
        assert request.url.path == "/places"
        return httpx.Response(200, json={"places": [{"website":"https://dealer.example", "cid":"1", "rating":4, "ratingCount":10, "reviews":[{"name":"Person"}], "phoneNumber":"private"}]})
    result = Serper("synthetic-test", httpx.Client(transport=httpx.MockTransport(respond))).places("Dealer", "City")
    assert set(result[0]) == {"website", "cid", "rating", "ratingCount"}


def test_chat_absent_html_is_unknown_and_widgets_are_classified():
    assert detect_chat("", "https://dealer.example").value is None
    assert detect_chat('<script src="https://embed.tawk.to/abc"></script>', "https://dealer.example").value == "Tawk"
    assert detect_chat('<a href="https://wa.me/123">Chat</a>', "https://dealer.example").value == "WhatsApp"
    assert detect_chat("<h1>Dealer</h1>", "https://dealer.example").value == "Niet gedetecteerd"
    assert classify_vacancy("Serviceadviseur en klantenservice") == {"serviceadviseur", "klantenservice"}


def test_vacancy_requires_function_page_and_ignores_closed_job():
    class Search:
        def places(self, *args): return []
    class Fetcher:
        def fetch(self, url):
            if url.endswith("vacatures"):
                return url, "Klantenservice", '<a href="/jobs/1">Serviceadviseur</a><a href="/jobs/2">Receptionist</a>'
            if url.endswith("1"):
                return url, "Serviceadviseur. Solliciteer nu", "<h1>Serviceadviseur</h1>"
            return url, "Receptionist. Vacature is vervuld", ""
    row = lead()
    LeadSearch(None, None, Search(), Fetcher())._sam_signals(row, '<a href="/vacatures">Werken bij</a>')
    assert row.signals["vacancies"].value == "serviceadviseur"
    assert row.signals["vacancies"].source.endswith("/jobs/1")
    assert "https://dealer.example/jobs/1" in row.sources


def test_export_has_points_signals_and_scorecard_explanation():
    row = lead(rating=3.7, reviews=3000, vacancies="receptie", chat="Niet gedetecteerd")
    book = load_workbook(BytesIO(export_xlsx(brief(), [row], context_name="SAM", created_at="2026-10-08", expires_at="2027-01-06")))
    headers = [cell.value for cell in book.active[1]]
    assert {"Rating", "Reviews", "Vacatures", "Chattool", "Punten", "Tier"}.issubset(headers)
    assert book.active.cell(2, headers.index("Punten") + 1).value == 70
    assert any("82" in str(cell.value) for line in book["Over deze lijst"] for cell in line)
